import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import sqlite3
import datetime
import db

def generate_excel_workbook(db_path=None):
    conn = db.get_db_connection()
    c = conn.cursor()

    wb = openpyxl.Workbook()
    
    # Sheet 1: Card view
    ws_cards = wb.active
    ws_cards.title = "Карта Расселения"
    ws_cards.views.sheetView[0].showGridLines = True

    # Sheet 2: Flat List
    ws_list = wb.create_sheet(title="Полный Список Жильцов")
    ws_list.views.sheetView[0].showGridLines = True

    # Sheet 3: Duty Roster
    ws_duty = wb.create_sheet(title="График Дежурств")
    ws_duty.views.sheetView[0].showGridLines = True

    # --- 1. Export Flat List Sheet ---
    headers = ["ID", "№", "ФИО Проживающего", "Никнейм", "Ссылка на профиль", "Пол", "Комната", "Этаж", "Статус", "Дата создания"]
    ws_list.append(headers)
    
    c.execute("""
    SELECT r.id, r.num_code, r.full_name, r.nickname, r.profile_url, r.gender, r.room_number, rm.floor, r.status, r.created_at
    FROM residents r
    LEFT JOIN rooms rm ON r.room_number = rm.room_number
    WHERE r.status != 'evicted'
    ORDER BY rm.floor ASC, r.room_number ASC, r.full_name ASC
    """)
    rows = c.fetchall()
    
    for row in rows:
        status_ru = {
            "permanent": "Постоянный",
            "14_days": "14 дней (Временный)",
            "waiting": "В очереди",
            "evicted": "Выселен"
        }.get(row["status"], row["status"])
        
        ws_list.append([
            row["id"],
            row["num_code"] or "—",
            row["full_name"],
            row["nickname"] or "—",
            row["profile_url"] or "—",
            "Мужской" if row["gender"] == "M" else "Женский",
            row["room_number"] or "Без комнаты",
            row["floor"] or "—",
            status_ru,
            row["created_at"]
        ])

    font_th = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    fill_th = PatternFill(start_color="1E3A8A", fill_type="solid")
    for cell in ws_list[1]:
        cell.font = font_th
        cell.fill = fill_th
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws_list.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_list.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # --- 2. Visual Card View Sheet ---
    ws_cards.merge_cells("A1:K1")
    title_cell = ws_cards["A1"]
    title_cell.value = "   АКТУАЛЬНАЯ КАРТА РАССЕЛЕНИЯ ОБЩЕЖИТИЯ (ИЗ БАЗЫ ДАННЫХ)"
    title_cell.font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E3A8A", fill_type="solid")
    title_cell.alignment = Alignment(vertical="center", horizontal="left")
    ws_cards.row_dimensions[1].height = 40

    c.execute("SELECT COUNT(*) as total FROM residents WHERE status != 'evicted'")
    total_res = c.fetchone()["total"]
    c.execute("SELECT COUNT(*) as f2 FROM residents WHERE room_number LIKE '2%' AND status != 'evicted'")
    f2_res = c.fetchone()["f2"]
    c.execute("SELECT COUNT(*) as f7 FROM residents WHERE room_number LIKE '7%' AND status != 'evicted'")
    f7_res = c.fetchone()["f7"]

    font_kpi_num = Font(name="Segoe UI", size=16, bold=True, color="1E3A8A")
    font_kpi_label = Font(name="Segoe UI", size=9, color="64748B")
    fill_kpi = PatternFill(start_color="F1F5F9", fill_type="solid")

    for lbl_col, val_col, label, val in [("B3:C3", "B4:C4", "Всего жильцов", total_res), ("D3:E3", "D4:E4", "2 Этаж (Девушки)", f2_res), ("F3:G3", "F4:G4", "7 Этаж (Парни)", f7_res)]:
        ws_cards.merge_cells(lbl_col); ws_cards.merge_cells(val_col)
        l_cell = ws_cards[lbl_col.split(":")[0]]
        l_cell.value = label; l_cell.font = font_kpi_label; l_cell.fill = fill_kpi; l_cell.alignment = Alignment(horizontal="center", vertical="center")
        v_cell = ws_cards[val_col.split(":")[0]]
        v_cell.value = val; v_cell.font = font_kpi_num; v_cell.fill = fill_kpi; v_cell.alignment = Alignment(horizontal="center", vertical="center")

    # --- 3. Duty Roster Sheet (Floor 7 Only) ---
    now = datetime.date.today()
    duty_data = db.get_duty_schedule_for_month(now.year, now.month, floor=7)

    ws_duty.merge_cells("A1:H1")
    duty_title = ws_duty["A1"]
    duty_title.value = f"   ГРАФИК ДЕЖУРСТВ ПО КУХНЕ (7 ЭТАЖ) — {now.strftime('%B %Y').upper()}"
    duty_title.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    duty_title.fill = PatternFill(start_color="1E3A8A", fill_type="solid")
    duty_title.alignment = Alignment(vertical="center", horizontal="left")
    ws_duty.row_dimensions[1].height = 36

    duty_headers = ["Дата", "День", "Этаж", "Дежурная Комната", "Жильцы комнаты", "Ответственный (Лидер)", "Статус", "Заметки"]
    ws_duty.append([])
    ws_duty.append(duty_headers)

    for cell in ws_duty[3]:
        cell.font = font_th
        cell.fill = fill_th
        cell.alignment = Alignment(horizontal="center", vertical="center")

    items = duty_data.get(7, [])
    for item in items:
        res_names = ", ".join(r["full_name"] for r in item["residents"])
        primary_name = item["primary_resident"]["full_name"] if item.get("primary_resident") else "—"
        status_text = {
            "pending": "Запланировано",
            "completed": "Выполнено",
            "skipped": "Пропущено",
            "replaced": "Заменено"
        }.get(item["status"], item["status"])

        ws_duty.append([
            item["date"],
            item["day_of_week"],
            "7 Этаж",
            f"Комната {item['room_number']}",
            res_names or "—",
            primary_name,
            status_text,
            item["notes"] or ""
        ])

    for col in ws_duty.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_duty.column_dimensions[col_letter].width = max(max_len + 3, 12)

    conn.close()
    return wb

def export_db_to_excel(db_path, target_excel_path):
    wb = generate_excel_workbook(db_path)
    try:
        wb.save(target_excel_path)
        return True, "Файл успешно обновлен"
    except PermissionError:
        fallback_path = target_excel_path.replace(".xlsx", "_обновленный.xlsx")
        try:
            wb.save(fallback_path)
            return True, f"Основной файл открыт в Excel. Сохранено в '{os.path.basename(fallback_path)}'"
        except Exception as e:
            return False, str(e)
    except Exception as e:
        return False, str(e)
